"""Check the paper snapshot, Word/PDF content, figures, and archive integrity."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from pypdf import PdfReader


HERE = Path(__file__).resolve().parent
manifest = json.loads((HERE / "paper_data" / "manifest.json").read_text(encoding="utf-8"))
figures = [item for item in manifest["items"] if item["category"] == "figure"]
data = [item for item in manifest["items"] if item["category"] == "data"]
assert len(figures) == 10 and len(data) == 36
for item in manifest["items"]:
    target = HERE / item["paper_path"]
    source = HERE.parent.parent / item["source_path"]
    assert target.is_file() and source.is_file(), item["paper_path"]
    digest = hashlib.sha256(target.read_bytes()).hexdigest()
    assert digest == item["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()

md = (HERE / "manuscript_full_v2.md").read_text(encoding="utf-8")
doc = Document(HERE / "manuscript_full_v2.docx")
doc_text = "\n".join(p.text for p in doc.paragraphs)
released_pdf = HERE / "manuscript_full_v2.pdf"
pdf_path = released_pdf if released_pdf.is_file() else HERE / "qa_render" / "manuscript_full_v2.pdf"
qa_pdf = HERE / "qa_render" / "manuscript_full_v2.pdf"
if released_pdf.is_file() and qa_pdf.is_file():
    assert hashlib.sha256(released_pdf.read_bytes()).digest() == hashlib.sha256(qa_pdf.read_bytes()).digest(), "Stale public/QA PDF pair"
pdf = PdfReader(str(pdf_path))
pdf_text = "\n".join(page.extract_text() or "" for page in pdf.pages)
assert len(doc.inline_shapes) == 10
assert len(doc.tables) == 8  # Seven main tables and supplementary data index.
assert len(pdf.pages) >= 13
for text in (md, doc_text, pdf_text):
    text = re.sub(r"\s+", " ", text)
    for phrase in ("Runtime Image Replacement", "2,475 ordered", "199,699", "100 Mbps",
                   "Figure 7", "Figure 8", "Figure 9", "Figure 10", "Table 5", "Table 6", "Table 7",
                   "4,308", "350,185", "214,982", "20/20", "bad-FCS", "biological"):
        assert phrase in text, phrase
assert "45AEAAAE" in md and "6C233D31" in md
archive = HERE / "CNS2FPGA_paper_package_v2.zip"
names = set()
if archive.is_file():
    with zipfile.ZipFile(archive) as package:
        assert package.testzip() is None
        names = set(package.namelist())
        for name in ("manuscript_full_v2.md", "manuscript_full_v2.docx",
                     "qa_render/manuscript_full_v2.pdf", "paper_data/manifest.json",
                     "figures/figure7_ethernet_reconfiguration.png",
                     "figures/figure8_p0_repeated_board_validation.png",
                     "figures/figure9_long_ordered_events.png",
                     "figures/figure10_100m_board_validation.png",
                     "paper_data/ethernet_deployment_evidence_v1.json",
                     "paper_data/old_new_board_comparison_v1.json",
                     "paper_data/p0_board_audit_v1.json",
                     "paper_data/long_event_board_audit_v1.json",
                     "paper_data/expanded_bit_100m_audit_v1.json"):
            assert name in names, name
        for item in manifest["items"]:
            assert item["paper_path"] in names
print(json.dumps({"status": "PASS", "figures": len(figures), "data": len(data),
                  "docx_images": len(doc.inline_shapes), "docx_tables": len(doc.tables),
                  "pdf_pages": len(pdf.pages), "archive_members": len(names)}, indent=2))
