"""Build the self-contained paper figures, data snapshot, and editable DOCX.

Source data are copied unchanged. The manifest records both the original path
and SHA-256 of each copied file, so the paper can be audited against the runs.
"""

from __future__ import annotations

import hashlib
import json
import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


HERE = Path(__file__).resolve().parent
CODE = HERE.parent
PROJECT = CODE.parent

FIGURES = {
    "figure1_pipeline.png": HERE / "figures" / "pipeline_paper.png",
    "figure2_courtship_response.png": CODE / "10_Complete Courtship-Song Circuit Experiment" / "outputs" / "courtship_song_v1" / "response_curve.png",
    "figure3_courtship_raster.png": CODE / "10_Complete Courtship-Song Circuit Experiment" / "outputs" / "courtship_song_v1" / "spike_raster.png",
    "figure4_robustness.png": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "degradation_curve.png",
    "figure5_quantization.png": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "quantization_envelope.png",
    "figure6_lateralized_baseline.png": CODE / "13_Baselines and Ablations" / "outputs" / "baseline_v0" / "lateralized_baseline.png",
    "figure7_ethernet_reconfiguration.png": HERE / "figures" / "figure7_ethernet_reconfiguration.png",
}

DATA = {
    "courtship_response_comparison.csv": CODE / "10_Complete Courtship-Song Circuit Experiment" / "outputs" / "courtship_song_v1" / "response_comparison.csv",
    "courtship_latency_throughput.csv": CODE / "10_Complete Courtship-Song Circuit Experiment" / "outputs" / "courtship_song_v1" / "latency_throughput.csv",
    "courtship_resource_timing.json": CODE / "10_Complete Courtship-Song Circuit Experiment" / "outputs" / "courtship_song_v1" / "resource_timing.json",
    "courtship_short_spike_raster.csv": CODE / "10_Complete Courtship-Song Circuit Experiment" / "outputs" / "courtship_song_v1" / "short_spike_raster.csv",
    "courtship_degradation_summary.csv": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "degradation_summary.csv",
    "courtship_condition_metrics.csv": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "condition_metrics.csv",
    "courtship_fpga_noise_comparison.csv": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "fpga_noise_comparison.csv",
    "courtship_fpga_noise_per_step_verification.json": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "fpga_noise_per_step_verification.json",
    "courtship_quantization_envelope.csv": CODE / "11_Courtship-Song Robustness" / "outputs" / "courtship_song_robustness_v2" / "quantization_envelope.csv",
    "visual_board_verification.json": CODE / "12_Visual-to-Steering Circuit" / "outputs" / "visual_left_board_v0" / "verification.json",
    "visual_board_events.csv": CODE / "12_Visual-to-Steering Circuit" / "outputs" / "visual_left_board_v0" / "parsed" / "events.csv",
    "visual_board_counts.csv": CODE / "12_Visual-to-Steering Circuit" / "outputs" / "visual_left_board_v0" / "parsed" / "counts.csv",
    "visual_manual_vs_automatic.csv": CODE / "13_Baselines and Ablations" / "outputs" / "baseline_v0" / "manual_vs_automatic.csv",
    "visual_structural_comparison.json": CODE / "13_Baselines and Ablations" / "outputs" / "baseline_v0" / "structural_comparison.json",
    "visual_width_ablation.csv": CODE / "13_Baselines and Ablations" / "outputs" / "baseline_v0" / "width_ablation.csv",
    "ethernet_deployment_evidence_v1.json": CODE / "15_Ethernet_Runtime_Deployment" / "reports" / "deployment_evidence_v1.json",
    "old_new_board_comparison_v1.json": CODE / "15_Ethernet_Runtime_Deployment" / "reports" / "old_new_board_comparison_v1.json",
}

ALT_TEXT = {
    "figure1_pipeline.png": "Flowchart from MaleCNS source through extraction and compilation to CPU and FPGA verification",
    "figure2_courtship_response.png": "Four courtship population response curves over nine input pulse intervals; float, fixed, and FPGA series overlap",
    "figure3_courtship_raster.png": "Three identical 350-ms spike rasters for float CPU, fixed CPU, and FPGA",
    "figure4_robustness.png": "Four panels showing pC1 count deviation under synapse deletion, neuron failure, weight jitter, and input noise",
    "figure5_quantization.png": "Bar chart showing only the narrowest fixed-point profile exceeds the five percent response-error criterion",
    "figure6_lateralized_baseline.png": "Left and right unilateral LC10a input produce ipsilateral DNa02 spikes in the extracted graph but none in the simplified graph",
    "figure7_ethernet_reconfiguration.png": "One AXKU115 bitstream accepts visual and courtship graph images over UDP; the courtship smoke8 FPGA counts exactly overlap fixed CPU at every step",
}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def copy_assets() -> None:
    figures_dir = HERE / "figures"
    data_dir = HERE / "paper_data"
    figures_dir.mkdir(exist_ok=True)
    data_dir.mkdir(exist_ok=True)
    manifest = {"schema": "cns2fpga.paper_data.v2", "source_version": "MaleCNS v1.0", "items": []}
    for category, mapping, destination in (
        ("figure", FIGURES, figures_dir),
        ("data", DATA, data_dir),
    ):
        for name, source in mapping.items():
            if not source.is_file():
                raise FileNotFoundError(source)
            target = destination / name
            if source.resolve() != target.resolve():
                shutil.copy2(source, target)
            manifest["items"].append({
                "category": category,
                "paper_path": str(target.relative_to(HERE)).replace("\\", "/"),
                "source_path": str(source.relative_to(PROJECT)).replace("\\", "/"),
                "sha256": digest(target),
                "bytes": target.stat().st_size,
            })
    (data_dir / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")


def set_font(style, name: str, size: float, *, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    edge = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        edge.append(el)
    tc_pr.append(edge)


def cell_margin(cell, top=85, left=110, bottom=85, right=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        tag = OxmlElement(f"w:{side}")
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")
        mar.append(tag)
    tc_pr.append(mar)


def clean_inline(value: str) -> str:
    value = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", value)
    return value.replace("`", "").replace("*", "")


def add_text(doc, value: str, style: str | None = None):
    para = doc.add_paragraph(style=style)
    para.add_run(clean_inline(value))
    return para


def add_table(doc, rows: list[list[str]]) -> None:
    count = len(rows[0])
    table = doc.add_table(rows=0, cols=count)
    table.autofit = False
    table.style = "Table Grid"
    widths = ([1.2, 1.65, 1.65, 1.65] if count == 4 else
              [1.2, 1.05, 1.05, 1.05, 1.7] if count == 5 else
              [1.15, 0.75, 0.9, 1.2, 1.35, 0.9] if count == 6 else
              [2.7, 2.0, 2.0] if count == 3 else
              [2.2, 4.5] if count == 2 else [6.7 / count] * count)
    # Keep every table within the 6.8-inch text block.
    scale = 6.7 / sum(widths)
    widths = [item * scale for item in widths]
    for idx, values in enumerate(rows):
        row = table.add_row()
        if idx == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(values):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)
            borders(cell)
            if idx == 0:
                shade(cell, "DDE9EF")
            elif idx % 2 == 0:
                shade(cell, "F5F8FA")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = idx < len(rows) - 1
            run = p.add_run(clean_inline(value))
            run.font.name = "Aptos"
            run.font.size = Pt(8.3)
            run.font.bold = idx == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_figure(doc, relpath: str) -> None:
    path = HERE / relpath
    with Image.open(path) as im:
        width_px, height_px = im.size
    width = 6.55
    height = width * height_px / width_px
    if height > 4.8:
        height = 4.8
        width = height * width_px / height_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    shape = p.add_run().add_picture(str(path), width=Inches(width), height=Inches(height))
    shape._inline.docPr.set("descr", ALT_TEXT[path.name])


def create_docx() -> Path:
    lines = (HERE / "manuscript_full_v2.md").read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    set_font(doc.styles["Normal"], "Aptos", 10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.14
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, before, after in (("Title", 17.5, 0, 10), ("Heading 1", 13.5, 15, 5), ("Heading 2", 11.5, 10, 4)):
        style = doc.styles[name]
        set_font(style, "Aptos Display" if name == "Title" else "Aptos", size, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    # Word's built-in Title style may carry a blue bottom rule even after
    # the font color is changed. Remove that inherited paragraph border.
    title_ppr = doc.styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    caption_style = doc.styles.add_style("Paper Caption", 1)
    set_font(caption_style, "Aptos", 9)
    caption_style.paragraph_format.space_before = Pt(2)
    caption_style.paragraph_format.space_after = Pt(9)
    doc.core_properties.title = clean_inline(lines[0][2:])
    doc.core_properties.subject = "Connectome-constrained FPGA computation and verification"
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            add_text(doc, line[2:], "Title")
        elif line.startswith("## "):
            add_text(doc, line[3:], "Heading 1")
        elif line.startswith("### "):
            add_text(doc, line[4:], "Heading 2")
        elif line.startswith("!["):
            match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if not match:
                raise ValueError(line)
            add_figure(doc, match.group(1))
        elif line.startswith("*Figure "):
            add_text(doc, line.strip("*"), "Paper Caption")
        elif line.startswith("Table ") and ". " in line:
            p = add_text(doc, line, "Paper Caption")
            p.paragraph_format.keep_with_next = True
        elif line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [part.strip() for part in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", part) for part in values):
                    rows.append(values)
                i += 1
            add_table(doc, rows)
            continue
        else:
            add_text(doc, line)
        i += 1
    out = HERE / "manuscript_full_v2.docx"
    doc.save(out)
    return out


def create_archive() -> Path:
    manifest = json.loads((HERE / "paper_data" / "manifest.json").read_text(encoding="utf-8"))
    members = [HERE / "manuscript_full_v2.md", HERE / "manuscript_full_v2.docx",
               HERE / "paper_data" / "manifest.json"]
    pdf = HERE / "qa_render" / "manuscript_full_v2.pdf"
    if pdf.is_file():
        members.append(pdf)
    members.extend(HERE / item["paper_path"] for item in manifest["items"])
    out = HERE / "CNS2FPGA_paper_package_v2.zip"
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for member in members:
            archive.write(member, arcname=str(member.relative_to(HERE)).replace("\\", "/"))
    return out


if __name__ == "__main__":
    copy_assets()
    output = create_docx()
    print(output)
    print(create_archive())
